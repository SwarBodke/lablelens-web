"""Generate truthful PDF/DOCX/JSON screening reports from stored inspections."""
from __future__ import annotations

import io
import json
from datetime import datetime, timezone
from typing import Any, Dict

from ruleset import metadata as ruleset_metadata


def _safe(v: Any) -> str:
    if v is None:
        return "Not detected"
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False, indent=2)
    return str(v)


def _ocr_label(record: Dict[str, Any]) -> str:
    if record.get("source_type") == "ecommerce" and record.get("ocr_confidence") is None:
        return "N/A - direct listing text"
    value = record.get("ocr_confidence")
    if value is None:
        return "N/A"
    try:
        return f"{float(value):.1f}%"
    except Exception:
        return str(value)


def report_payload(record: Dict[str, Any]) -> Dict[str, Any]:
    return {
        "report_generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "report_type": "Legal Metrology screening aid - not an enforcement order",
        "ruleset": ruleset_metadata(),
        "inspection": record,
    }


def _audit_rows(record: Dict[str, Any]):
    record_hash = record.get("record_hash") or (record.get("audit") or {}).get("record_hash")
    previous_hash = record.get("previous_hash") or (record.get("audit") or {}).get("previous_hash")
    if not record_hash:
        return []
    return [
        ["Audit algorithm", "SHA-256 per-session record chain"],
        ["Record hash", record_hash],
        ["Previous record hash", previous_hash or "GENESIS"],
    ]


def build_pdf(record: Dict[str, Any]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, rightMargin=34, leftMargin=34, topMargin=32, bottomMargin=32)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Small", parent=styles["BodyText"], fontSize=8.5, leading=11))
    story = [
        Paragraph("LableLens - Legal Metrology Screening Report", styles["Title"]),
        Paragraph(f"Inspection INS-{record.get('id', '-')}", styles["Heading2"]),
        Paragraph(
            "Evidence-led screening under the Legal Metrology (Packaged Commodities) Rules, 2011. "
            "The six groups are a practical automation baseline and this report is not an enforcement order or a substitute for current Gazette text, exemptions or category-specific law.",
            styles["Small"],
        ),
        Spacer(1, 10),
    ]
    status = record.get("compliance_status") or record.get("status") or "needs_review"
    summary_rows = [
        ["Product", _safe(record.get("product_name"))],
        ["Status", str(status).replace("_", " ").title()],
        ["Source", str(record.get("source_type") or "physical").replace("_", " ").title()],
        ["Evidence scope", str(record.get("evidence_scope") or "partial").replace("_", " ").title()],
        ["Created", _safe(record.get("created_at"))],
        ["OCR confidence", _ocr_label(record)],
        ["Extraction completeness", f"{record.get('extraction_completeness', 0)}%"],
        ["Ruleset", record.get("ruleset_id") or ruleset_metadata()["id"]],
    ] + _audit_rows(record)
    t = Table(summary_rows, colWidths=[130, 380])
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#d8dee8")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f4f6f8")),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("LEADING", (0, 0), (-1, -1), 10.5),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.extend([t, Spacer(1, 14), Paragraph("Six core check groups", styles["Heading2"])])

    checks = record.get("checks") or {}
    if isinstance(checks, str):
        try:
            checks = json.loads(checks)
        except Exception:
            checks = {}
    for key, item in checks.items():
        if not isinstance(item, dict):
            continue
        title = item.get("title") or key.replace("_", " ").title()
        st = item.get("status", "needs_review").replace("_", " ").title()
        story.append(Paragraph(f"{title} - {st}", styles["Heading3"]))
        story.append(Paragraph(_safe(item.get("summary") or item.get("evidence") or ""), styles["Small"]))
        if item.get("rule"):
            story.append(Paragraph(f"Reference: {_safe(item.get('rule'))}", styles["Small"]))
        for issue in item.get("issues") or []:
            escaped = str(issue).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(f"- {escaped}", styles["Small"]))
        story.append(Spacer(1, 5))

    story.extend([Spacer(1, 10), Paragraph("Transcription / listing evidence", styles["Heading2"])])
    raw = str(record.get("raw_text") or "No transcription stored.")[:12000]
    raw = raw.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>")
    story.append(Paragraph(raw, styles["Small"]))

    evidence = record.get("evidence") or []
    if isinstance(evidence, str):
        try:
            evidence = json.loads(evidence)
        except Exception:
            evidence = []
    if evidence:
        story.extend([Spacer(1, 10), Paragraph("Image evidence hashes", styles["Heading2"])])
        for ev in evidence:
            story.append(Paragraph(f"SHA-256: {_safe(ev.get('sha256'))} ({ev.get('bytes', 0)} bytes)", styles["Small"]))

    doc.build(story)
    return out.getvalue()


def build_docx(record: Dict[str, Any]) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("LableLens - Legal Metrology Screening Report", level=0)
    doc.add_paragraph(f"Inspection INS-{record.get('id', '-')}")
    doc.add_paragraph(
        "Evidence-led screening under the Legal Metrology (Packaged Commodities) Rules, 2011. "
        "This is a screening aid, not an enforcement order or a substitute for current Gazette text and applicable category-specific law."
    )
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    status = record.get("compliance_status") or record.get("status") or "needs_review"
    summary = [
        ("Product", record.get("product_name")),
        ("Status", status.replace("_", " ").title()),
        ("Source", (record.get("source_type") or "physical").replace("_", " ").title()),
        ("Evidence scope", (record.get("evidence_scope") or "partial").replace("_", " ").title()),
        ("Created", record.get("created_at")),
        ("OCR confidence", _ocr_label(record)),
        ("Extraction completeness", f"{record.get('extraction_completeness', 0)}%"),
        ("Ruleset", record.get("ruleset_id") or ruleset_metadata()["id"]),
    ]
    for k, v in summary:
        cells = table.add_row().cells
        cells[0].text = str(k)
        cells[1].text = _safe(v)

    audit = _audit_rows(record)
    if audit:
        doc.add_heading("Audit chain", level=1)
        for key, value in audit:
            doc.add_paragraph(f"{key}: {value}")

    doc.add_heading("Six core check groups", level=1)
    checks = record.get("checks") or {}
    if isinstance(checks, str):
        try:
            checks = json.loads(checks)
        except Exception:
            checks = {}
    for key, item in checks.items():
        if not isinstance(item, dict):
            continue
        doc.add_heading(
            f"{item.get('title') or key} - {item.get('status', 'needs_review').replace('_', ' ').title()}",
            level=2,
        )
        doc.add_paragraph(_safe(item.get("summary") or item.get("evidence") or ""))
        if item.get("rule"):
            doc.add_paragraph(f"Reference: {item.get('rule')}")
        for issue in item.get("issues") or []:
            doc.add_paragraph(str(issue), style="List Bullet")

    doc.add_heading("Transcription / listing evidence", level=1)
    doc.add_paragraph(str(record.get("raw_text") or "No transcription stored.")[:12000])
    evidence = record.get("evidence") or []
    if isinstance(evidence, str):
        try:
            evidence = json.loads(evidence)
        except Exception:
            evidence = []
    if evidence:
        doc.add_heading("Image evidence hashes", level=1)
        for ev in evidence:
            doc.add_paragraph(f"SHA-256: {ev.get('sha256')} ({ev.get('bytes', 0)} bytes)", style="List Bullet")

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
